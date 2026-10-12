"""Word(.docx)出力。

Sample/ にある議事メモテンプレートと同じ書式（Meiryo UI、表構成、見出し）を再現する。
本文はタイムスタンプを除いた逐語テキストをそのまま段落として貼り付ける（要約・整形はしない）。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Pt

from .merge import extract_body_paragraphs

FONT_NAME = "Meiryo UI"
DOCX_VERSION_SUFFIX = "_Ver0.1"
BULLET_FONT_NAME = "Wingdings"
BODY_SIZE_PT = 10.5
TABLE_WIDTH_DXA = 9645
LABEL_COL_DXA = 1081
VALUE_COL_DXA = 8564

# Sample/(サンプル)20261007_収支DB_議事メモ_v0.1.docx の word/numbering.xml
# abstractNumId=27（numId=29）から実測した全9階層（ilvl 0-8）。
# Word上でTabキーによりレベルを下げた際、サンプルと同じ記号が循環して表示されるよう
# 全レベルを複製する（■→➢→✧→●→➢→✧→●→➢→✧）。
BULLET_LEVELS = [
    {"ilvl": 0, "char": "", "left": 440, "hanging": 440, "tentative": False},
    {"ilvl": 1, "char": "", "left": 880, "hanging": 440, "tentative": False},
    {"ilvl": 2, "char": "", "left": 1320, "hanging": 440, "tentative": False},
    {"ilvl": 3, "char": "", "left": 1760, "hanging": 440, "tentative": False},
    {"ilvl": 4, "char": "", "left": 2200, "hanging": 440, "tentative": True},
    {"ilvl": 5, "char": "", "left": 2640, "hanging": 440, "tentative": True},
    {"ilvl": 6, "char": "", "left": 3080, "hanging": 440, "tentative": True},
    {"ilvl": 7, "char": "", "left": 3520, "hanging": 440, "tentative": True},
    {"ilvl": 8, "char": "", "left": 3960, "hanging": 440, "tentative": True},
]


def _add_bullet_numbering(doc: Document) -> int:
    """Sampleと同じ9階層（■→➢→✧→●...循環）の箇条書き定義を追加し、使用するnumIdを返す。"""
    numbering_element = doc.part.numbering_part.element

    existing_abstract_ids = [
        int(e.get(qn("w:abstractNumId")))
        for e in numbering_element.findall(qn("w:abstractNum"))
    ]
    existing_num_ids = [
        int(e.get(qn("w:numId")))
        for e in numbering_element.findall(qn("w:num"))
    ]
    abstract_num_id = (max(existing_abstract_ids) + 1) if existing_abstract_ids else 0
    num_id = (max(existing_num_ids) + 1) if existing_num_ids else 1

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_num_id))

    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "hybridMultilevel")
    abstract_num.append(multi_level)

    for level in BULLET_LEVELS:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level["ilvl"]))
        if level["tentative"]:
            lvl.set(qn("w:tentative"), "1")

        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)

        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet")
        lvl.append(num_fmt)

        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), level["char"])
        lvl.append(lvl_text)

        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)

        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(level["left"]))
        ind.set(qn("w:hanging"), str(level["hanging"]))
        ppr.append(ind)
        lvl.append(ppr)

        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), BULLET_FONT_NAME)
        rfonts.set(qn("w:hAnsi"), BULLET_FONT_NAME)
        rfonts.set(qn("w:hint"), "default")
        rpr.append(rfonts)
        lvl.append(rpr)

        abstract_num.append(lvl)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id_ref = OxmlElement("w:abstractNumId")
    abstract_num_id_ref.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract_num_id_ref)

    # スキーマ順序: すべての abstractNum の後にすべての num が続く必要がある。
    existing_nums = numbering_element.findall(qn("w:num"))
    if existing_nums:
        existing_nums[0].addprevious(abstract_num)
    else:
        numbering_element.append(abstract_num)
    numbering_element.append(num)

    return num_id


def _set_bullet_level(paragraph, num_id: int, ilvl: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numid_el = OxmlElement("w:numId")
    numid_el.set(qn("w:val"), str(num_id))
    numpr.append(ilvl_el)
    numpr.append(numid_el)
    ppr.append(numpr)


def _add_field_run(paragraph, field_code: str, cached_text: str) -> None:
    """PAGE / NUMPAGES 等のWordフィールドを段落に追加する（ページ数に応じて自動更新される）。"""
    run_begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld_begin)

    run_instr = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    run_instr._r.append(instr)

    run_sep = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    run_cached = paragraph.add_run(cached_text)
    _set_run_font(run_cached)

    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def _add_page_number_footer(doc: Document) -> None:
    """右下に「P{現在ページ}/{総ページ数}」を表示するフッターを追加する（Sampleのfooter1.xmlを再現）。"""
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run_p = paragraph.add_run("P")
    _set_run_font(run_p)
    _add_field_run(paragraph, "PAGE", "1")
    run_slash = paragraph.add_run("/")
    _set_run_font(run_slash)
    _add_field_run(paragraph, "NUMPAGES", "1")


def _set_run_font(run, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(BODY_SIZE_PT)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:eastAsia", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)


def _set_cell_width(cell, width_dxa: int) -> None:
    cell.width = Pt(width_dxa / 20)  # dxa -> pt (1pt = 20dxa)
    tcPr = cell._element.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = tcPr.makeelement(qn("w:tcW"), {})
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")


def _shade_cell(cell, fill_hex: str) -> None:
    tcPr = cell._element.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill_hex,
    })
    tcPr.append(shd)


def _write_cell_text(cell, text: str, *, center: bool = False, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    _set_run_font(run, bold=bold)


def build_document(title: str, body_paragraphs: list[str]) -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_SIZE_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    # python-docxの既定テンプレートは段落後10pt・行間1.15倍が入っているが、
    # Sampleにはその指定が無い（≒0pt・単一行間隔）。Sampleに合わせて明示的にゼロにする。
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    _add_page_number_footer(doc)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    title_row = table.rows[0]
    title_row.cells[0].merge(title_row.cells[1])
    _write_cell_text(title_row.cells[0], title, center=True)
    _shade_cell(title_row.cells[0], "C0C0C0")
    _set_cell_width(title_row.cells[0], TABLE_WIDTH_DXA)

    field_rows = [
        ("日時", ""),
        ("場所", ""),
        ("参加者\n(敬称略)", ""),
        ("資料", ""),
    ]
    for i, (label, value) in enumerate(field_rows, start=1):
        row = table.rows[i]
        _write_cell_text(row.cells[0], label, center=True)
        _set_cell_width(row.cells[0], LABEL_COL_DXA)
        _write_cell_text(row.cells[1], value)
        _set_cell_width(row.cells[1], VALUE_COL_DXA)

    doc.add_paragraph()

    heading = doc.add_paragraph()
    run = heading.add_run("【会議内容】")
    _set_run_font(run)

    num_id = _add_bullet_numbering(doc)

    topic = doc.add_paragraph()
    _set_bullet_level(topic, num_id, ilvl=0)
    run = topic.add_run("XXX")
    _set_run_font(run, bold=True)
    run.font.underline = True

    for para_text in body_paragraphs:
        p = doc.add_paragraph()
        _set_bullet_level(p, num_id, ilvl=1)
        run = p.add_run(para_text)
        _set_run_font(run)

    return doc


def write_docx(output_dir: Path, name: str, transcript_text: str) -> Path:
    body_paragraphs = extract_body_paragraphs(transcript_text)
    doc = build_document(title=name, body_paragraphs=body_paragraphs)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{name}{DOCX_VERSION_SUFFIX}.docx"
    doc.save(path)
    return path
